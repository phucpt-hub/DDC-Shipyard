# Ghi file Excel "Nhật ký làm sạch dữ liệu": tổng hợp quy tắc + chi tiết từng ô đã chuẩn hóa.
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

RED = '7A1712'
FILL_H = PatternFill('solid', fgColor=RED)
FILL = {'Đã chuẩn hóa': PatternFill('solid', fgColor='E6F3EA'), 'Quy ước': PatternFill('solid', fgColor='EFEAE9'), 'Cần xác nhận': PatternFill('solid', fgColor='FBF0DC')}
THIN = Side(style='thin', color='E2DCDB')
F = 'Arial'

def _header(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(row=row, column=i, value=c)
        x.font = Font(name=F, bold=True, color='FFFFFF', size=10); x.fill = FILL_H
        x.alignment = Alignment(vertical='center', wrap_text=True); x.border = Border(bottom=THIN)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30

def ghi(log, path, ngay):
    df = pd.DataFrame(log)
    wb = Workbook(); ws = wb.active; ws.title = 'TONG_HOP'
    ws['A1'] = 'NHẬT KÝ LÀM SẠCH DỮ LIỆU – CHUỖI THÉP SC VŨNG TÀU'; ws['A1'].font = Font(name=F, bold=True, size=14, color=RED)
    ws['A2'] = f'Số liệu ngày {ngay}. Báo cáo đã tính theo dữ liệu sau khi làm sạch; file gốc không bị sửa. Sheet CHI_TIET liệt kê từng ô để các phòng sửa file gốc.'
    ws['A2'].font = Font(name=F, italic=True, size=10)
    cols = ['Mã', 'Quy tắc', 'Loại', 'File', 'Số dòng', 'KL liên quan (tấn)']
    _header(ws, 4, cols, [7, 90, 15, 30, 10, 16])
    r = 5
    if len(df):
        g = df.groupby(['ma', 'quy_tac', 'nhom', 'file'], sort=False).agg(n=('row', 'size'), kl=('kl', 'sum')).reset_index()
        for x in g.itertuples():
            vals = [x.ma, x.quy_tac, x.nhom, x.file, int(x.n), round(x.kl / 1000, 3)]
            for i, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=i, value=v); c.font = Font(name=F, size=10); c.border = Border(bottom=THIN)
                c.alignment = Alignment(vertical='top', wrap_text=(i == 2))
                if i == 3: c.fill = FILL.get(x.nhom)
                if i == 6: c.number_format = '#,##0.0'
                if i == 5: c.number_format = '#,##0'
            r += 1
        c = ws.cell(row=r, column=1, value='Tổng'); c.font = Font(name=F, bold=True)
        ws.cell(row=r, column=5, value=f'=SUM(E5:E{r-1})').number_format = '#,##0'
        ws.cell(row=r, column=6, value=f'=SUM(F5:F{r-1})').number_format = '#,##0.0'
        for i in (5, 6): ws.cell(row=r, column=i).font = Font(name=F, bold=True)
    r += 2
    for t in ['Đã chuẩn hóa: báo cáo đã tự sửa theo quy tắc, nên sửa luôn trong file gốc cho lần sau.',
              'Quy ước: giữ theo file, ghi lại để biết cách tính.',
              'Cần xác nhận: báo cáo giữ nguyên như file, nhờ phòng phụ trách kiểm tra và sửa file gốc nếu sai.']:
        ws.cell(row=r, column=2, value=t).font = Font(name=F, size=10); r += 1
    ws.freeze_panes = 'A5'
    # chi tiết
    wd = wb.create_sheet('CHI_TIET')
    cols = ['Mã', 'Loại', 'File', 'Sheet', 'Dòng Excel', 'LSX / Mã / Nesting', 'Dự án', 'Cột', 'Giá trị cũ', 'Giá trị mới (báo cáo dùng)', 'KL (kg)', 'Quy tắc']
    _header(wd, 1, cols, [7, 14, 26, 18, 10, 30, 34, 24, 28, 24, 11, 70])
    if len(df):
        df = df.sort_values(['nhom', 'ma', 'file', 'sheet', 'row'], key=lambda s: s.map({'Cần xác nhận': 0, 'Đã chuẩn hóa': 1, 'Quy ước': 2}) if s.name == 'nhom' else s)
        for j, x in enumerate(df.itertuples(), 2):
            vals = [x.ma, x.nhom, x.file, x.sheet, x.row, x.obj, x.du_an, x.cot, x.cu, x.moi, round(x.kl, 2), x.quy_tac]
            for i, v in enumerate(vals, 1):
                c = wd.cell(row=j, column=i, value=v); c.font = Font(name=F, size=9.5)
                if i == 2: c.fill = FILL.get(x.nhom)
                if i == 11: c.number_format = '#,##0.00'
    wd.freeze_panes = 'A2'; wd.auto_filter.ref = f'A1:L{max(2, len(df) + 1)}'
    wb.save(path)
    return len(df)
