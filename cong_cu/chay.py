# Chạy toàn bộ: đọc Excel trong du_lieu/ -> phân tích -> index.html + lich_su/<ngày>/
import os, sys, re, io, glob, json, pickle, subprocess, datetime, shutil, runpy
import pandas as pd

CC = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(CC)
DL = os.path.join(ROOT, 'du_lieu')
TMP = os.path.join(CC, '_tmp'); os.makedirs(TMP, exist_ok=True)
EXT = ('.xlsx', '.xlsm', '.xls', '.xlsb')

def log(*a): print('>>', *a, flush=True)

def commit_time(p):
    try:
        t = subprocess.run(['git', '-C', ROOT, 'log', '-1', '--format=%ct', '--', p], capture_output=True, text=True).stdout.strip()
        return int(t) if t else os.path.getmtime(p)
    except Exception:
        return os.path.getmtime(p)

def name_date(p):
    # ngày trong tên file: '... 01.10.2026', 'TỒN KHO 28.09', '30-09-26'
    now = datetime.datetime.utcnow() + datetime.timedelta(hours=7)
    for m in re.finditer(r'(?<!\d)(\d{1,2})[._-](\d{1,2})(?:[._-](\d{2,4}))?(?!\d)', os.path.basename(p)):
        y = int(m.group(3)) if m.group(3) else now.year
        y = y + 2000 if y < 100 else y
        try: return datetime.date(y, int(m.group(2)), int(m.group(1)))
        except ValueError: continue
    return None

def commit_date(p):
    return datetime.datetime.fromtimestamp(commit_time(p), datetime.timezone(datetime.timedelta(hours=7))).date()

def files_in(folder, theo_ngay_tai_len=False):
    fs = [f for f in glob.glob(os.path.join(DL, folder, '*')) if f.lower().endswith(EXT) and not os.path.basename(f).startswith('~$')]
    # cùng tên nhưng khác đuôi (.xlsx và .xlsb) → ưu tiên .xlsx
    best = {}
    for f in fs:
        stem, ext = os.path.splitext(os.path.basename(f)); rank = EXT.index(ext.lower())
        if stem not in best or rank < best[stem][0]: best[stem] = (rank, f)
    skipped = sorted(set(fs) - {v[1] for v in best.values()})
    for f in skipped: log('Bỏ qua (trùng tên, đã có bản .xlsx):', os.path.basename(f))
    # file mới nhất xếp cuối: theo ngày trong tên file, rồi theo thời điểm tải lên
    # theo_ngay_tai_len: tên file không có ngày thì lấy ngày tải file lên (dùng cho file tồn kho cập nhật tay)
    return sorted([v[1] for v in best.values()], key=lambda f: (name_date(f) or (commit_date(f) if theo_ngay_tai_len else datetime.date(1900, 1, 1)), commit_time(f)))

def _cell(v):
    if v is None: return 'NaN'
    if isinstance(v, float):
        if v != v: return 'NaN'
        return str(int(v)) if v.is_integer() and abs(v) < 1e15 else ('%.10g' % v)
    if isinstance(v, bool): return str(v)
    if isinstance(v, int): return str(v)
    if isinstance(v, (pd.Timestamp, datetime.datetime)): return pd.Timestamp(v).strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(v, datetime.date): return v.strftime('%Y-%m-%d 00:00:00')
    if isinstance(v, datetime.time): return v.strftime('%H:%M:%S')
    s = ' '.join(str(v).split())            # gộp xuống dòng như bản markdown
    return s.replace('_', '\\_').replace('*', '\\*') if s else 'NaN'

# Chỉ đọc các sheet cần cho báo cáo (bỏ BCSL, CHI PHÍ, danh sach NV, biểu đồ... cho nhanh)
SHEETS_CUTTING = ['list chi tiết', '2-list chi tiết', 'lệnh sx', 'kế hoạch', 'máy móc']
SHEETS_NESTING = ['list chi tiet', 'phieu vat tu', 'hao hut vat tu']

def sheet_names(xlsx):
    if xlsx.lower().endswith('.xlsb'):
        from pyxlsb import open_workbook
        with open_workbook(xlsx) as wb: return list(wb.sheets)
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True); n = wb.sheetnames; wb.close(); return n

def to_sheets(xlsx, out_pkl, only=None):
    """Excel -> bảng chuỗi theo từng sheet, cùng định dạng với bản .md đã dùng để phân tích.
    only: danh sách tên sheet cần đọc (không phân biệt hoa thường, bỏ khoảng trắng đầu/cuối); None = đọc hết."""
    names = sheet_names(xlsx)
    pick = [n for n in names if only is None or n.strip().lower() in only]
    if only is not None: log('  bỏ qua', len(names) - len(pick), 'sheet không dùng:', ', '.join(n for n in names if n not in pick)[:160])
    raw = pd.read_excel(xlsx, sheet_name=pick, header=0, dtype=object, engine='pyxlsb' if xlsx.lower().endswith('.xlsb') else None)
    out = {}
    for k, df in raw.items():
        hdr = [_cell(c) if not str(c).startswith('Unnamed') else str(c) for c in df.columns]
        n = len(hdr)
        data = [[_cell(v) for v in row] for row in df.itertuples(index=False, name=None)]
        d2 = pd.DataFrame(data, columns=[f'c{i}' for i in range(n)]); d2.attrs['hdr'] = hdr; out[k.strip() if k.strip() != k and k.strip() in ('TỔNG HỢP NHẬP (TỒN + XUẤT)',) else k] = d2
    pickle.dump(out, open(out_pkl, 'wb'))
    log('Đọc', os.path.basename(xlsx), '→', len(out), 'sheet:', ', '.join(list(out)[:8]))
    return out

def is_html(f):
    """File .xls xuất từ phần mềm thường là trang HTML đổi đuôi (không phải Excel thật)."""
    with open(f, 'rb') as fh: h = fh.read(2048).lstrip().lower()
    return h.startswith(b'<') or b'<html' in h or b'<table' in h

def html_sheets(f):
    """Đọc file HTML đổi đuôi .xls → 1 bảng (bảng có nhiều dòng nhất), cùng định dạng với to_sheets."""
    raw = open(f, 'rb').read()
    for enc in ('utf-8', 'utf-16', 'cp1258'):
        try: txt = raw.decode(enc); break
        except UnicodeDecodeError: continue
    else: txt = raw.decode('utf-8', 'replace')
    df = max(pd.read_html(io.StringIO(txt)), key=len)
    if isinstance(df.columns, pd.MultiIndex): df.columns = [str(c[-1]) for c in df.columns]
    hdr = [_cell(c) if not str(c).startswith('Unnamed') else str(c) for c in df.columns]
    d2 = pd.DataFrame([[_cell(v) for v in row] for row in df.itertuples(index=False, name=None)], columns=[f'c{i}' for i in range(len(hdr))]); d2.attrs['hdr'] = hdr
    return {'Sheet1': d2}

def kho_sets():
    """Các bộ dữ liệu tồn kho, cũ → mới. Một bộ = 1 file đặt thẳng trong 2_TON_KHO,
    hoặc 1 thư mục con (vd 2_TON_KHO/25.10.2026) chứa 1–3 file Nhập / Xuất / Tồn. Ngày = ngày trong tên (thư mục/file), không có thì ngày tải lên."""
    base = os.path.join(DL, '2_TON_KHO'); sets = []
    for f in files_in('2_TON_KHO', theo_ngay_tai_len=True):
        sets.append((name_date(f) or commit_date(f), commit_time(f), [f], os.path.splitext(os.path.basename(f))[0]))
    for sub in sorted(glob.glob(os.path.join(base, '*'))):
        if not os.path.isdir(sub): continue
        fs = sorted(f for f in glob.glob(os.path.join(sub, '*')) if f.lower().endswith(EXT) and not os.path.basename(f).startswith('~$'))
        if fs: sets.append((name_date(sub) or max(name_date(f) or commit_date(f) for f in fs), max(commit_time(f) for f in fs), fs, 'TỒN KHO ' + os.path.basename(sub)))
    return sorted(sets, key=lambda s: (s[0], s[1]))

def doc_kho(files, out_pkl):
    """Đọc bộ file tồn kho vào 1 pickle. File chỉ có 1 sheet thì đặt tên sheet = tên file
    (để nhận ra file Nhập / Xuất / Tồn theo chữ trong tên file)."""
    out = {}
    for f in files:
        stem = os.path.splitext(os.path.basename(f))[0]
        if is_html(f): sh = html_sheets(f); log('Đọc', os.path.basename(f), '(HTML) →', len(sh['Sheet1']), 'dòng')
        else: sh = to_sheets(f, os.path.join(TMP, 'kho_tmp.pkl'))
        if len(sh) == 1: out[stem] = list(sh.values())[0]
        else:
            for k, v in sh.items(): out[k if k not in out else stem + ' · ' + k] = v
    pickle.dump(out, open(out_pkl, 'wb'))
    return out

def tom_tat(D):
    """Số liệu chính của một kỳ báo cáo (để so sánh giữa các kỳ)."""
    P = D.get('projects', [])
    open_lsx = [l for p in P for l in p.get('lsx', []) if (l.get('pc') or 0) < 0.999 and (l.get('kl') or 0) > 0]
    def buy(p):
        n = p.get('nes')
        if n:
            return sum(a['t'] for a in n.get('alloc', []) if a.get('k') == 'Phải mua')
        return (p.get('kho') or {}).get('buy')
    return {'date': D.get('date'), 'tot': D.get('tot', {}), 'bal': D.get('bal', {}),
            'nes_tp': sum(((p.get('nes') or {}).get('flow') or {}).get('tp', 0) for p in P),
            'open_lsx': len(open_lsx), 'open_t': sum(l['kl'] * (1 - (l.get('pc') or 0)) for l in open_lsx),
            'proj': {p['id']: {'kl': p.get('kl'), 'req': p.get('req'), 'cut': p.get('cut'), 'bg': p.get('bg'), 'buy': buy(p),
                               'need': (p.get('kho') or {}).get('need')} for p in P}}

def ky_truoc(ngay):
    """Tìm báo cáo đã lưu gần nhất TRƯỚC ngày báo cáo, lấy số liệu để so sánh."""
    ls = os.path.join(ROOT, 'lich_su')
    if not os.path.isdir(ls): return None
    for d in sorted([x for x in os.listdir(ls) if re.match(r'\d{4}-\d{2}-\d{2}$', x) and x < ngay], reverse=True):
        try:
            h = open(os.path.join(ls, d, 'index.html'), encoding='utf-8').read()
            m = re.search(r'const D=(\{.*?\});\s*\nconst P=', h, re.S)
            if m: return tom_tat(json.loads(m.group(1).replace('<\\/', '</')))
        except Exception as e:
            log('Không đọc được báo cáo cũ', d, '-', e)
    return None

def doc_viec(cfg):
    """Việc cần xử lý / cần quyết định: sheet VIEC_CAN_XU_LY (file cấu hình hoặc file VIEC_LINK tải từ link)."""
    out = []
    for f in sorted(cfg, key=lambda f: not os.path.basename(f).upper().startswith('VIEC_LINK')):
        try:
            names = sheet_names(f)
            sh = next((n for n in names if n.strip().upper().replace(' ', '_') in ('VIEC_CAN_XU_LY', 'VIỆC_CẦN_XỬ_LÝ')), None)
            if sh is None and os.path.basename(f).upper().startswith('VIEC_LINK'): sh = names[0]
            if sh is None: continue
            df = pd.read_excel(f, sheet_name=sh, dtype=object)
        except Exception as e:
            log('Không đọc được sheet việc cần xử lý trong', os.path.basename(f), '-', e); continue
        cols = {str(c).strip().lower(): c for c in df.columns}
        def col(*keys):
            for k in keys:
                for c in cols:
                    if k in c: return cols[c]
        cv, cd, cp, ch, cs, cg = col('việc', 'viec', 'nội dung'), col('dự án', 'du an'), col('phụ trách', 'phu trach', 'đơn vị'), col('hạn', 'deadline', 'han'), col('trạng thái', 'trang thai'), col('ghi chú', 'ghi chu')
        if cv is None: continue
        for _, r in df.iterrows():
            v = r.get(cv)
            if v is None or (isinstance(v, float) and v != v) or not str(v).strip(): continue
            def s(c):
                x = r.get(c) if c is not None else None
                if x is None or (isinstance(x, float) and x != x): return ''
                if isinstance(x, (pd.Timestamp, datetime.datetime, datetime.date)): return pd.Timestamp(x).strftime('%Y-%m-%d')
                return str(x).strip()
            out.append({'viec': str(v).strip(), 'du_an': s(cd), 'phu_trach': s(cp), 'han': s(ch), 'tt': s(cs), 'ghi_chu': s(cg)})
        if out: break
    return out

def main():
    cut = files_in('1_CUTTING'); kho = kho_sets(); nes = files_in('3_NESTING'); cfg = files_in('4_CAU_HINH')
    if not cut: sys.exit('THIẾU file Cutting trong du_lieu/1_CUTTING')
    if not kho: sys.exit('THIẾU file tồn kho trong du_lieu/2_TON_KHO')
    f_cut, (d_kho, _, f_khos, ten_kho) = cut[-1], kho[-1]
    log('File Cutting:', os.path.basename(f_cut)); log('File tồn kho:', ten_kho, '←', ', '.join(os.path.basename(f) for f in f_khos))
    # Ngày báo cáo = ngày mới nhất trong tên file Cutting / tồn kho; không có ngày trong tên thì lấy hôm nay (giờ VN)
    now = datetime.datetime.utcnow() + datetime.timedelta(hours=7)
    d_cut = name_date(f_cut)
    ds = [d for d in (d_cut, d_kho) if d and d <= now.date()]
    ngay = max(ds) if ds else now.date()
    os.environ.update(NGAY_CUTTING=str(d_cut or ''), NGAY_KHO=str(d_kho or ''))
    if d_cut and d_kho and d_cut != d_kho: log(f'Lưu ý: file Cutting ngày {d_cut:%d/%m}, file tồn kho ngày {d_kho:%d/%m} – phần kho tính theo ngày {d_kho:%d/%m}')
    log('Ngày báo cáo:', ngay)
    to_sheets(f_cut, os.path.join(CC, 'sheets.pkl'), SHEETS_CUTTING)
    doc_kho(f_khos, os.path.join(CC, 'kho', 'sheets.pkl'))
    nfiles = []
    NES_OK = []
    for i, f in enumerate(nes):
        p = os.path.join(TMP, f'nes{i}.pkl')
        try:
            sh = to_sheets(f, p, SHEETS_NESTING)
        except Exception as e:
            log('KHÔNG đọc được file nesting', os.path.basename(f), '-', e); continue
        if not any(k.strip().lower() == 'phieu vat tu' for k in sh):
            log('Bỏ qua', os.path.basename(f), '- không có sheet "Phieu vat tu" (không phải file nesting)'); continue
        nfiles.append([p, os.path.splitext(os.path.basename(f))[0]])
    os.environ.update(NGAY_BAO_CAO=str(ngay), NESTING_FILES=json.dumps(nfiles, ensure_ascii=False),
                      FILE_CUTTING=os.path.splitext(os.path.basename(f_cut))[0], FILE_KHO=ten_kho,
                      OUT_JSON=os.path.join(TMP, 'data.json'), OUT_LOG=os.path.join(TMP, 'clean_log.pkl'),
                      OUT_CT=os.path.join(ROOT, 'chi_tiet'))
    cfg_map = [f for f in cfg if not os.path.basename(f).upper().startswith('VIEC_LINK')]
    if cfg_map: os.environ['FILE_CAU_HINH'] = cfg_map[-1]
    os.chdir(CC)
    runpy.run_path('data.py', run_name='__main__')
    D = json.load(open(os.path.join(TMP, 'data.json'), encoding='utf-8'))
    D['prev'] = ky_truoc(str(ngay))
    D['viec'] = doc_viec(cfg)
    log('So với kỳ trước:', D['prev']['date'] if D['prev'] else 'không có', '· Việc cần xử lý:', len(D['viec']), 'dòng')
    data = json.dumps(D, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
    html = open('template.html', encoding='utf-8').read().replace('/*DATA*/null', data)
    open(os.path.join(ROOT, 'index.html'), 'w', encoding='utf-8').write(html)
    hd = os.path.join(ROOT, 'lich_su', str(ngay)); os.makedirs(hd, exist_ok=True)
    open(os.path.join(hd, 'index.html'), 'w', encoding='utf-8').write(html.replace('href="lich_su/"', 'href="../"'))
    # trang danh sách báo cáo cũ
    days = sorted([d for d in os.listdir(os.path.join(ROOT, 'lich_su')) if re.match(r'\d{4}-\d{2}-\d{2}$', d)], reverse=True)
    items = ''.join(f'<li><a href="{d}/">Báo cáo ngày {d[8:10]}/{d[5:7]}/{d[:4]}</a></li>' for d in days)
    open(os.path.join(ROOT, 'lich_su', 'index.html'), 'w', encoding='utf-8').write(
        '<!doctype html><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>Báo cáo cũ</title>'
        '<style>body{font:15px/1.6 Arial,sans-serif;max-width:640px;margin:0 auto;padding:24px 16px;color:#18232d;background:#f3f5f6}a{color:#1d5f86}</style>'
        f'<h1>Lịch sử báo cáo chuỗi thép</h1><p><a href="../">← Báo cáo mới nhất</a></p><ul>{items}</ul>')
    # công cụ làm sạch: nhật ký Excel
    sys.path.insert(0, CC); import nhat_ky
    n = nhat_ky.ghi(pickle.load(open(os.path.join(TMP, 'clean_log.pkl'), 'rb')), os.path.join(ROOT, 'nhat_ky_lam_sach.xlsx'), ngay.strftime('%d/%m/%Y'))
    shutil.copy(os.path.join(ROOT, 'nhat_ky_lam_sach.xlsx'), os.path.join(hd, 'nhat_ky_lam_sach.xlsx'))
    log('Nhật ký làm sạch:', n, 'ô')
    log('XONG: index.html + lich_su/' + str(ngay))

if __name__ == '__main__':
    main()
